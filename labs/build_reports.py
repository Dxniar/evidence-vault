"""Build four Word laboratory reports from the maintained Markdown sources.

Run with the bundled workspace Python after the document operation marker.
This helper does not run the marker, install tools, render, or publish files.
Use --check for a read-only source/image preflight before the authoring step.
The parent workflow must render and inspect every page before delivery.
"""
from __future__ import annotations

import argparse
import math
import os
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import unquote, urlsplit

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
LABS = ROOT / "labs"
BODY_FONT = "Times New Roman"
CODE_FONT = "Courier New"
BODY_SIZE = 11.5
BODY_WIDTH_INCHES = 6.5
IMAGE_MAX_HEIGHT_INCHES = 6.4
REPORTS = {
    1: ("ЛР_1_Двухфакторный_вход.docx", "Лабораторная работа 1 Двухфакторный вход Evidence Vault"),
    2: ("ЛР_2_Федеративный_вход.docx", "Лабораторная работа 2 Федеративный вход Evidence Vault"),
    3: ("ЛР_3_Журналирование.docx", "Лабораторная работа 3 Журналирование Evidence Vault"),
    4: ("ЛР_4_Политика_Windows.docx", "Лабораторная работа 4 Политика Windows для Evidence Vault"),
}


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    rows: list[list[str]] | None = None
    ordered: bool = False
    ordinal: int = 1
    image_path: Path | None = None


def require_bundled_python():
    expected = (Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime"
                / "dependencies" / "python" / "python.exe")
    if os.name == "nt" and os.path.normcase(str(Path(sys.executable).resolve())) != os.path.normcase(str(expected.resolve())):
        raise RuntimeError(f"Запустите скрипт через Python бандла workspace dependencies: {expected}")


def heading_words(value: str) -> str:
    value = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", value)
    return " ".join(re.sub(r"[\W_]+", " ", value, flags=re.UNICODE).split())


def table_cells(line: str) -> list[str]:
    value = line.strip().strip("|")
    # Keep an escaped vertical bar inside its cell instead of treating it as a separator.
    return [cell.strip().replace(r"\|", "|") for cell in re.split(r"(?<!\\)\|", value)]


def table_separator(line: str) -> bool:
    cells = table_cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def resolve_image(target: str, source: Path) -> Path:
    target = target.strip()
    if target.startswith("<") and target.endswith(">"):
        target = target[1:-1]
    target = unquote(target)
    if re.match(r"^[a-zA-Z]:[\\/]", target):
        path = Path(target)
    else:
        parsed = urlsplit(target)
        if parsed.scheme and parsed.scheme != "file":
            raise ValueError(f"Изображение должно быть локальным файлом: {target}")
        if parsed.scheme == "file":
            path = Path(unquote(parsed.path.lstrip("/") if os.name == "nt" else parsed.path))
        else:
            path = Path(target)
    if not path.is_absolute():
        direct = source.parent / path
        project = ROOT / path
        path = direct if direct.is_file() else project
    path = path.resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Не найдено изображение из {source.name}: {path}")
    with Image.open(path) as picture:
        picture.verify()
    return path


IMAGE_PATTERN = re.compile(r'^!\[([^\]]*)\]\((.+?)(?:\s+"[^\"]*")?\)\s*$')
LIST_PATTERN = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.+)$")


def parse_markdown(source: Path) -> list[Block]:
    lines = source.read_text(encoding="utf-8-sig").splitlines()
    result = []
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("```"):
            code = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            if index >= len(lines):
                raise ValueError(f"Незакрытый блок кода: {source}")
            result.append(Block("code", "\n".join(code)))
            index += 1
            continue
        header = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if header:
            result.append(Block("heading", header.group(2), level=len(header.group(1))))
            index += 1
            continue
        image = IMAGE_PATTERN.match(stripped)
        if image:
            result.append(Block("image", image.group(1), image_path=resolve_image(image.group(2), source)))
            index += 1
            continue
        if index + 1 < len(lines) and "|" in line and table_separator(lines[index + 1]):
            rows = [table_cells(line)]
            index += 2
            while index < len(lines) and lines[index].strip() and "|" in lines[index]:
                rows.append(table_cells(lines[index]))
                index += 1
            columns = len(rows[0])
            if any(len(row) != columns for row in rows):
                raise ValueError(f"Разное число столбцов в таблице: {source}")
            result.append(Block("table", rows=rows))
            continue
        item = LIST_PATTERN.match(line)
        if item:
            indent, marker, body = item.groups()
            index += 1
            continuation = []
            while index < len(lines) and lines[index].strip() and re.match(r"^\s{2,}\S", lines[index]) and not LIST_PATTERN.match(lines[index]):
                continuation.append(lines[index].strip())
                index += 1
            body = " ".join([body, *continuation])
            result.append(Block("list", body, level=len(indent.expandtabs(4)) // 2,
                                ordered=marker[0].isdigit(),
                                ordinal=int(re.match(r"\d+", marker).group()) if marker[0].isdigit() else 1))
            continue
        if re.fullmatch(r"[-*_]{3,}", stripped):
            index += 1
            continue
        paragraph = [stripped]
        index += 1
        while index < len(lines) and lines[index].strip():
            candidate = lines[index]
            if (candidate.strip().startswith(("#", "```")) or LIST_PATTERN.match(candidate)
                    or IMAGE_PATTERN.match(candidate.strip())
                    or (index + 1 < len(lines) and "|" in candidate and table_separator(lines[index + 1]))):
                break
            paragraph.append(candidate.strip())
            index += 1
        result.append(Block("paragraph", " ".join(paragraph)))
    return result


def set_font(target, name, size):
    target.font.name = name
    target.font.size = Pt(size)
    target.font.color.rgb = RGBColor(0, 0, 0)
    if hasattr(target, "_element"):
        element = target._element
        rpr = element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + key), name)
        for theme in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"):
            fonts.attrib.pop(qn("w:" + theme), None)
        color = rpr.find(qn("w:color"))
        if color is not None:
            for attr in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn("w:" + attr), None)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.35)
    normal = doc.styles["Normal"]
    set_font(normal, BODY_FONT, BODY_SIZE)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 18), ("Subtitle", 12), ("Heading 1", 15),
                       ("Heading 2", 13), ("Heading 3", 12)):
        style = doc.styles[name]
        set_font(style, BODY_FONT, size)
        style.font.bold = name.startswith("Heading")
        style.paragraph_format.space_before = Pt(8 if name.startswith("Heading") else 0)
        style.paragraph_format.space_after = Pt(5 if name.startswith("Heading") else 12)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.line_spacing = 1.05
        ppr = style.element.get_or_add_pPr()
        for border in list(ppr.findall(qn("w:pBdr"))):
            ppr.remove(border)
    for name in ("List Bullet", "List Number", "Caption"):
        style = doc.styles[name]
        set_font(style, BODY_FONT, BODY_SIZE if name != "Caption" else 10)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.1
    if "Lab Code" not in doc.styles:
        code = doc.styles.add_style("Lab Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = doc.styles["Lab Code"]
    code.base_style = normal
    set_font(code, CODE_FONT, 9)
    code.paragraph_format.line_spacing = 1
    code.paragraph_format.space_before = Pt(5)
    code.paragraph_format.space_after = Pt(9)
    code.paragraph_format.left_indent = Inches(0.08)
    code.paragraph_format.right_indent = Inches(0.08)
    code.paragraph_format.keep_together = False
    code.paragraph_format.keep_with_next = False
    code.paragraph_format.widow_control = True
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.subject = "Лабораторная работа на примере Evidence Vault"
    doc.core_properties.keywords = "Evidence Vault, лабораторная работа"
    # Explicit Russian proofing, without changing text or placeholders.
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        rdefault = defaults.find(qn("w:rPrDefault"))
        if rdefault is not None:
            rpr = rdefault.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                rdefault.append(rpr)
            language = rpr.find(qn("w:lang"))
            if language is None:
                language = OxmlElement("w:lang")
                rpr.append(language)
            language.set(qn("w:val"), "ru-RU")


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*.+?\*\*|\*[^*]+\*|\[[^\]]+\]\((?:<[^>]+>|[^)]+)\))")


def add_hyperlink(paragraph, label, address, *, bold=False, italic=False, size=BODY_SIZE):
    address = address.strip()
    if address.startswith("<") and address.endswith(">"):
        address = address[1:-1]
    if re.match(r"(?:javascript|data):", address, flags=re.IGNORECASE):
        raise ValueError("Недопустимая схема гиперссылки")
    link = OxmlElement("w:hyperlink")
    if address.startswith("#"):
        link.set(qn("w:anchor"), address[1:])
    else:
        link.set(qn("r:id"), paragraph.part.relate_to(address, RT.HYPERLINK, is_external=True))
    link.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attribute), BODY_FONT)
    props.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    point_size = OxmlElement("w:sz")
    point_size.set(qn("w:val"), str(round(size * 2)))
    props.append(point_size)
    for active, tag in ((bold, "b"), (italic, "i")):
        if active:
            props.append(OxmlElement("w:" + tag))
    run.append(props)
    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = re.sub(r"[`*_]", "", label)
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def inline(paragraph, text, *, bold=False, italic=False, size=BODY_SIZE):
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            add_run(paragraph, text[position:match.start()], bold=bold, italic=italic, size=size)
        value = match.group(0)
        if value.startswith("`"):
            run = add_run(paragraph, value[1:-1], bold=bold, italic=italic, size=min(size, 10))
            set_font(run, CODE_FONT, min(size, 10))
        elif value.startswith("**"):
            inline(paragraph, value[2:-2], bold=True, italic=italic, size=size)
        elif value.startswith("*"):
            inline(paragraph, value[1:-1], bold=bold, italic=True, size=size)
        else:
            label, address = re.fullmatch(r"\[([^\]]+)\]\((.+)\)", value).groups()
            add_hyperlink(paragraph, label, address, bold=bold, italic=italic, size=size)
        position = match.end()
    if position < len(text):
        add_run(paragraph, text[position:], bold=bold, italic=italic, size=size)


def add_run(paragraph, text, *, bold=False, italic=False, size=BODY_SIZE):
    run = paragraph.add_run(text)
    set_font(run, BODY_FONT, size)
    run.bold = bold
    run.italic = italic
    return run


def plain_text(value):
    value = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", value)
    return re.sub(r"[`*_]", "", value)


def column_widths(rows):
    count = len(rows[0])
    if count == 1:
        return [BODY_WIDTH_INCHES]
    scores = []
    for index in range(count):
        texts = [plain_text(row[index]) for row in rows]
        typical = sum(min(160, len(text)) for text in texts) / max(1, len(texts))
        longest = max(len(text) for text in texts)
        score = math.sqrt(max(12, typical * 0.7 + min(longest, 180) * 0.3))
        if all(len(text) <= 12 for text in texts):
            score *= 0.7
        scores.append(score)
    total = sum(scores)
    widths = [BODY_WIDTH_INCHES * score / total for score in scores]
    # Very short identifiers remain compact; prose columns retain useful width.
    minimum = min(0.75, BODY_WIDTH_INCHES / count * 0.75)
    for _ in range(3):
        deficits = sum(max(0, minimum - width) for width in widths)
        available = sum(max(0, width - minimum) for width in widths)
        if not deficits or not available:
            break
        widths = [minimum if width < minimum else width - deficits * (width - minimum) / available
                  for width in widths]
    scale = BODY_WIDTH_INCHES / sum(widths)
    return [width * scale for width in widths]


def add_table(doc, rows):
    keep_event_table = rows[0] == ["Поле", "Значение"] and len(rows) <= 12
    if keep_event_table:
        for paragraph in doc.paragraphs[-2:]:
            if paragraph.text.startswith(("Категория:", "Событие ")):
                paragraph.paragraph_format.keep_with_next = True
    widths = column_widths(rows)
    table = doc.add_table(rows=1, cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    props = table._tbl.tblPr
    borders = props.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    layout = props.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        props.append(layout)
    layout.set(qn("w:type"), "fixed")
    table_width = props.find(qn("w:tblW"))
    if table_width is not None:
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(round(BODY_WIDTH_INCHES * 1440)))
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for row_index, values in enumerate(rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for col_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge, amount in (("top", 90), ("bottom", 90), ("left", 100), ("right", 100)):
                element = OxmlElement("w:" + edge)
                element.set(qn("w:w"), str(amount))
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            tcpr.append(margins)
            if row_index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EDEDED")
                shade.set(qn("w:val"), "clear")
                tcpr.append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.paragraph_format.keep_with_next = keep_event_table and row_index < len(rows) - 1
            paragraph.paragraph_format.keep_together = False
            short_column = all(len(plain_text(item[col_index])) <= 16 for item in rows)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if short_column else WD_ALIGN_PARAGRAPH.LEFT
            inline(paragraph, value.replace("<br>", "\n").replace("<br/>", "\n"),
                   bold=row_index == 0, size=11)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 1
    spacer.add_run().font.size = Pt(3)


def add_image(doc, block, index):
    with Image.open(block.image_path) as picture:
        width_px, height_px = picture.size
    scale = min(BODY_WIDTH_INCHES / width_px, IMAGE_MAX_HEIGHT_INCHES / height_px)
    width = width_px * scale
    height = height_px * scale
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run()
    shape = run.add_picture(str(block.image_path), width=Inches(width), height=Inches(height))
    shape._inline.docPr.set("descr", plain_text(block.text) or f"Рисунок {index}")
    shape._inline.docPr.set("title", plain_text(block.text) or f"Рисунок {index}")
    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.keep_with_next = False
    inline(caption, f"Рисунок {index}. {block.text}" if block.text else f"Рисунок {index}", size=10)


def new_numbering(doc, start=1):
    """Create a Word list instance so independent procedure lists restart."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for name, value in (("start", str(start)), ("numFmt", "decimal"), ("lvlText", "%1."), ("lvlJc", "left")):
        child = OxmlElement("w:" + name)
        child.set(qn("w:val"), value)
        level.append(child)
    abstract.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)
    instance = OxmlElement("w:num")
    instance.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    instance.append(reference)
    numbering.append(instance)
    return num_id


def apply_numbering(paragraph, num_id):
    props = paragraph._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    reference = OxmlElement("w:numId")
    reference.set(qn("w:val"), str(num_id))
    numbering.append(level)
    numbering.append(reference)
    props.append(numbering)


def build(source, destination, title, blocks):
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = title
    heading = doc.add_paragraph(heading_words(title), style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    image_index = 0
    first_heading = True
    active_numbering = None
    for block in blocks:
        if block.kind != "list" or not block.ordered:
            active_numbering = None
        if block.kind == "heading":
            if first_heading and block.level == 1:
                first_heading = False
                continue
            first_heading = False
            level = min(3, max(1, block.level - 1))
            doc.add_paragraph(heading_words(block.text), style=f"Heading {level}")
        elif block.kind == "code":
            paragraph = doc.add_paragraph(style="Lab Code")
            run = paragraph.add_run(block.text)
            set_font(run, CODE_FONT, 9)
        elif block.kind == "table":
            add_table(doc, block.rows)
        elif block.kind == "image":
            image_index += 1
            add_image(doc, block, image_index)
        elif block.kind == "list":
            paragraph = doc.add_paragraph(style="List Number" if block.ordered else "List Bullet")
            if block.ordered:
                if active_numbering is None:
                    active_numbering = new_numbering(doc, block.ordinal)
                apply_numbering(paragraph, active_numbering)
            paragraph.paragraph_format.left_indent = Inches(0.2 + min(block.level, 3) * 0.2)
            paragraph.paragraph_format.first_line_indent = Inches(-0.15)
            inline(paragraph, block.text)
        else:
            paragraph = doc.add_paragraph()
            inline(paragraph, block.text)
            # Keep a short listing introduction with the immediately following code.
            if plain_text(block.text).startswith("Листинг"):
                paragraph.paragraph_format.keep_with_next = True
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    return {"source": str(source), "output": str(destination), "blocks": len(blocks), "images": image_index}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=LABS / "content")
    parser.add_argument("--output-dir", type=Path, default=LABS / "reports")
    parser.add_argument("--labs", type=int, nargs="+", choices=sorted(REPORTS), default=sorted(REPORTS))
    parser.add_argument("--check", action="store_true", help="Only parse sources and verify referenced local images")
    args = parser.parse_args()
    require_bundled_python()
    plans = []
    for number in args.labs:
        source = (args.source_dir / f"lab{number}.md").resolve()
        if not source.is_file():
            raise FileNotFoundError(f"Не найден источник лабораторной: {source}")
        blocks = parse_markdown(source)
        if not blocks:
            raise ValueError(f"Пустой отчёт: {source}")
        filename, title = REPORTS[number]
        plans.append((source, (args.output_dir / filename).resolve(), title, blocks))
    if args.check:
        for source, destination, title, blocks in plans:
            print(f"OK {source.name}: {len(blocks)} блоков; {sum(block.kind == 'image' for block in blocks)} изображений")
        print("Read-only preflight complete. No Word documents were created.")
        return 0
    # All four inputs/images are validated before the first document is written.
    for source, destination, title, blocks in plans:
        result = build(source, destination, title, blocks)
        print(f"Created {result['output']} ({result['blocks']} blocks, {result['images']} images)")
    print("Render and inspect all pages before delivery. This builder does not perform visual QA.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
